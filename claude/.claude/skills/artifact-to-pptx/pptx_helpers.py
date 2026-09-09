"""
pptx_helpers.py — artifact-to-pptx スキル用ヘルパ

HTML/Artifact を「全要素が編集可能な」PPTX に落とすための最小 API。
python-pptx の薄いラッパで，以下を helpers 側で強制する:

  - 16:9 (13.333 x 7.5 in) / blank レイアウト (slide_layouts[6])
  - 文字は必ず 15pt 以上 (MIN_PT)。15pt 未満の指定は自動で 15pt に引き上げ
  - 日本語フォントはメイリオ固定 (latin / ea / cs すべて Meiryo)
  - 影は消す (shadow.inherit = False)

座標・サイズはすべて「インチ (float)」で受ける。内部で EMU に変換する。

1 スライドの型:
    sl = slide(prs)
    header(sl, "タイトル", kicker="眉ラベル")
    card(sl, 0.6, 1.6, 5.8, 3.2, title="見出し",
         paras=[[("本文1", 15, False, INK)],
                [("強調", 15, True, ACCENT), ("＋通常", 15, False, INK)]])
    keyband(sl, "このスライドのキーメッセージ")

paras は必ず 2 重リスト:  [ 段落, 段落, ... ]  各段落 = [ run, run, ... ]
  run = (text, size, bold, color)
1 重で渡すと unpack エラーになる（意図的なガード）。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- 定数

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MIN_PT = 15.0            # これ未満の pt 指定は自動で引き上げる
JP_FONT = "Meiryo"      # 東アジアフォント固定

# single-accent パレット（濃紺アクセント / モノクロ寄り）。好みで差し替え可。
INK     = RGBColor(0x1A, 0x1A, 0x1A)   # 本文
MUTED   = RGBColor(0x6B, 0x72, 0x80)   # 眉ラベル・補足
ACCENT  = RGBColor(0x1F, 0x3A, 0x5F)   # アクセント（帯・強調）
ACCENT2 = RGBColor(0x3D, 0x6B, 0x9E)   # 補助アクセント（棒グラフ等）
LINE    = RGBColor(0xD0, 0xD5, 0xDB)   # 罫線
CARD_BG = RGBColor(0xF4, 0xF6, 0xF8)   # カード背景
BAND_FG = RGBColor(0xFF, 0xFF, 0xFF)   # 帯の文字
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

# タイポグラフィ既定
PT_TITLE   = 26.0
PT_KICKER  = 15.0
PT_BODY    = 15.0
PT_BAND    = 15.5

# レイアウト既定（インチ）
MARGIN_X   = 0.6
HEADER_Y   = 0.55
BAND_H     = 0.62


# ---------------------------------------------------------------- 低レベル

def new_prs():
    """16:9 の空プレゼンを作る。"""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def slide(prs):
    """blank レイアウト (index 6) のスライドを追加して返す。"""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _kill_shadow(shape):
    """python-pptx 既定のプリセット影を無効化する。"""
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def _set_dash(shape, val="dash"):
    """図形/線に点線 (a:prstDash) を注入する。"""
    ln = shape.line._get_or_add_ln()
    for old in ln.findall(qn("a:prstDash")):
        ln.remove(old)
    d = ln.makeelement(qn("a:prstDash"), {"val": val})
    ln.append(d)


def _apply_font(run, size, bold, color, font=JP_FONT):
    """run にフォント・サイズ・色を適用。size は MIN_PT で下限クランプ。
    latin だけでなく ea / cs も同フォントにして日本語を確実に当てる。"""
    size = max(float(size), MIN_PT)
    f = run.font
    f.size = Pt(size)
    f.bold = bool(bold)
    f.name = font                       # latin
    if color is not None:
        f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):        # 東アジア / complex-script も固定
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def _tf(shape, anchor="top", wrap=True, m=0.04):
    """text_frame を取り出して余白・折返し・縦位置を整える。"""
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(m)
    tf.margin_top = tf.margin_bottom = Inches(m)
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }[anchor]
    return tf


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
          "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


# ---------------------------------------------------------------- 図形

def box(sl, x, y, w, h, *, fill=None, line=None, line_w=1.0,
        dash=False, radius=False, shadow=False):
    """矩形（または角丸矩形）を描いて返す。テキストは持たない下地用。"""
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = sl.shapes.add_shape(shp_type, Inches(x), Inches(y),
                              Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            _set_dash(shp)
    if not shadow:
        _kill_shadow(shp)
    return shp


def txt(sl, text, x, y, w, h, *, size=PT_BODY, bold=False, color=INK,
        align="left", anchor="top", font=JP_FONT):
    """単一段落のテキストボックス。手早く 1 行〜数行を置く用。"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = _tf(tb, anchor=anchor)
    p = tf.paragraphs[0]
    p.alignment = _ALIGN[align]
    r = p.add_run()
    r.text = text
    _apply_font(r, size, bold, color, font)
    return tb


def _fill_paras(tf, paras, align="left", font=JP_FONT):
    """2 重リスト paras を text_frame に流し込む。
    paras = [ 段落, ... ] / 段落 = [ (text,size,bold,color), ... ]"""
    first = True
    for para in paras:                       # ← 1 重で渡すとここで unpack 失敗
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = _ALIGN[align]
        for text, size, bold, color in para:
            r = p.add_run()
            r.text = text
            _apply_font(r, size, bold, color, font)


def card(sl, x, y, w, h, *, title=None, paras=None, fill=CARD_BG,
         line=None, radius=True, title_size=16.0, align="left", pad=0.28):
    """下地の矩形 ＋ 任意タイトル ＋ 本文(paras)。最頻出の本文ブロック。"""
    box(sl, x, y, w, h, fill=fill, line=line, radius=radius)
    ty = y + pad
    if title is not None:
        tb = sl.shapes.add_textbox(Inches(x + pad), Inches(ty),
                                   Inches(w - 2 * pad), Inches(0.5))
        tf = _tf(tb)
        p = tf.paragraphs[0]
        p.alignment = _ALIGN[align]
        r = p.add_run()
        r.text = title
        _apply_font(r, title_size, True, ACCENT)
        ty += 0.52
    if paras:
        tb = sl.shapes.add_textbox(Inches(x + pad), Inches(ty),
                                   Inches(w - 2 * pad),
                                   Inches(max(0.3, y + h - ty - pad)))
        tf = _tf(tb)
        _fill_paras(tf, paras, align=align)
    return None


def mk_table(sl, x, y, w, rows, *, col_widths=None, header=True,
             font_size=PT_BODY, row_h=0.42, header_fill=ACCENT,
             header_fg=WHITE, body_fg=INK, zebra=CARD_BG):
    """rows(=文字列の 2 次元リスト)から編集可能な表を作る。
    大きすぎる表は呼び出し側で抜粋してから渡すこと（SKILL.md 参照）。"""
    nrow = len(rows)
    ncol = max(len(r) for r in rows)
    gtab = sl.shapes.add_table(nrow, ncol, Inches(x), Inches(y),
                               Inches(w), Inches(row_h * nrow))
    tbl = gtab.table
    tbl.first_row = header
    tbl.horz_banding = False
    if col_widths:
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            is_head = header and i == 0
            if is_head:
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
            elif zebra is not None and (i % 2 == 0):
                cell.fill.solid(); cell.fill.fore_color.rgb = zebra
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            para = cell.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = r[j] if j < len(r) else ""
            _apply_font(run, font_size, is_head,
                        header_fg if is_head else body_fg)
    return tbl


# ---------------------------------------------------------------- 型パーツ

def header(sl, title, *, kicker=None, y=HEADER_Y, x=MARGIN_X):
    """眉ラベル(kicker) ＋ タイトル(26pt) ＋ 下罫線。スライド上部の定型。"""
    w = SLIDE_W_IN - 2 * x
    cy = y
    if kicker:
        txt(sl, kicker, x, cy, w, 0.32, size=PT_KICKER, bold=True,
            color=MUTED)
        cy += 0.34
    txt(sl, title, x, cy, w, 0.7, size=PT_TITLE, bold=True, color=INK)
    cy += 0.74
    ln = box(sl, x, cy, w, 0.02, fill=ACCENT)     # 下罫線（細い塗り矩形）
    return cy + 0.06


def keyband(sl, message, *, bg=ACCENT, fg=BAND_FG, h=BAND_H):
    """スライド最下部のフル幅色帯 ＝ そのスライドのキーメッセージ。"""
    y = SLIDE_H_IN - h
    box(sl, 0, y, SLIDE_W_IN, h, fill=bg)
    tb = sl.shapes.add_textbox(Inches(MARGIN_X), Inches(y),
                               Inches(SLIDE_W_IN - 2 * MARGIN_X), Inches(h))
    tf = _tf(tb, anchor="middle")
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = message
    _apply_font(r, PT_BAND, True, fg)
    return tb


# ---------------------------------------------------------------- 図（shape 描画）

def vbars(sl, x, y, w, h, values, labels=None, *, color=ACCENT2,
          max_value=None, gap=0.3, baseline=True, value_fmt="{:.0f}"):
    """縦棒グラフを shape で描く（画像化しない＝Slides で編集可能）。
    values: 数値リスト / labels: 各棒の下ラベル / gap: 棒間の割合(0-1)。"""
    n = len(values)
    if n == 0:
        return
    mx = max_value if max_value is not None else max(values) or 1
    slot = w / n
    bw = slot * (1 - gap)
    off = (slot - bw) / 2
    axis_y = y + h
    for i, v in enumerate(values):
        bh = (float(v) / mx) * h if mx else 0
        bx = x + i * slot + off
        by = axis_y - bh
        box(sl, bx, by, bw, max(bh, 0.01), fill=color)
        txt(sl, value_fmt.format(v), bx - 0.1, by - 0.34, bw + 0.2, 0.3,
            size=MIN_PT, align="center", color=INK)
        if labels:
            txt(sl, str(labels[i]), bx - 0.1, axis_y + 0.05, bw + 0.2, 0.32,
                size=MIN_PT, align="center", color=MUTED)
    if baseline:
        box(sl, x, axis_y, w, 0.02, fill=LINE)


def funnel(sl, x, y, w, values, labels, *, step_h=0.62, gap=0.12,
           color=ACCENT):
    """ファネル（各段の幅 ∝ 値）を shape で描く。"""
    mx = max(values) or 1
    cy = y
    for v, lab in zip(values, labels):
        bw = w * (float(v) / mx)
        bx = x + (w - bw) / 2
        box(sl, bx, cy, bw, step_h, fill=color, radius=True)
        txt(sl, f"{lab}  {v}", x, cy, w, step_h, size=MIN_PT, bold=True,
            color=WHITE, align="center", anchor="middle")
        cy += step_h + gap


def gantt(sl, x, y, w, rows, *, unit_start=0, unit_end=10, row_h=0.42,
          gap=0.1, color=ACCENT2, grid=True):
    """簡易ガント。rows = [(label, start, end), ...]（start/end は単位量）。"""
    span = max(1e-6, unit_end - unit_start)
    label_w = 2.2
    chart_x = x + label_w
    chart_w = w - label_w
    cy = y
    if grid:
        ticks = 5
        for k in range(ticks + 1):
            gx = chart_x + chart_w * k / ticks
            box(sl, gx, y - 0.1, 0.008, row_h * len(rows) + 0.2, fill=LINE)
    for lab, s, e in rows:
        txt(sl, str(lab), x, cy, label_w - 0.1, row_h, size=MIN_PT,
            color=INK, anchor="middle")
        bx = chart_x + chart_w * (s - unit_start) / span
        bw = chart_w * (e - s) / span
        box(sl, bx, cy + 0.06, max(bw, 0.05), row_h - 0.12,
            fill=color, radius=True)
        cy += row_h + gap
